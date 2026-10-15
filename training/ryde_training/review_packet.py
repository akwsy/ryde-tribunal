"""Create validated review records from the owner-completed Excel review packet."""
import argparse
import json
from pathlib import Path

from .generate import ROLES, VERSION, outside_repository
from .review import DECISIONS, _approval_metadata, _target_json, _validate_standalone_judge


def _records(sheet, required_headers, label):
    values = list(sheet.iter_rows(values_only=True))
    if not values:
        raise ValueError(f"{label}: worksheet is empty")
    headers = [str(value or "").strip() for value in values[0]]
    missing = set(required_headers) - set(headers)
    if missing:
        raise ValueError(f"{label}: missing columns: {', '.join(sorted(missing))}")
    rows = []
    for values_row in values[1:]:
        if not any(value is not None and str(value).strip() for value in values_row):
            continue
        rows.append({header: values_row[index] if index < len(values_row) else None
                     for index, header in enumerate(headers) if header})
    return rows


def _text(value):
    return "" if value is None else str(value).strip()


def _decision_fields(record, prefix):
    decision = _text(record.get("Decision")).lower()
    if decision not in DECISIONS:
        raise ValueError(f"{prefix}: decision must be Pending, Approve, Correct, or Reject")
    meta = {
        "decision": decision,
        "reviewer": _text(record.get("Reviewer")) or None,
        "reviewed_at": _text(record.get("Reviewed at (ISO 8601)")) or None,
        "policy_version": _text(record.get("Policy version")) or None,
        "rationale": _text(record.get("Rationale")) or None,
    }
    if decision in {"approve", "correct"}:
        _approval_metadata(meta, prefix)
    return decision, meta


def apply_review_data(source_rows, ruling_rows, agent_rows):
    by_id = {row["sample"]["sample_id"]: row for row in source_rows}
    if len(by_id) != len(source_rows):
        raise ValueError("Review queue contains duplicate sample IDs")
    rulings = {}
    for record in ruling_rows:
        sample_id = _text(record.get("Sample ID"))
        if sample_id not in by_id:
            raise ValueError(f"Unknown sample ID in ruling sheet: {sample_id}")
        if sample_id in rulings:
            raise ValueError(f"Duplicate ruling review for {sample_id}")
        rulings[sample_id] = record
    if set(rulings) != set(by_id):
        raise ValueError("Ruling sheet must contain exactly one row for every queued case")

    agents = {}
    for record in agent_rows:
        sample_id = _text(record.get("Sample ID"))
        role = _text(record.get("Agent role"))
        key = (sample_id, role)
        if sample_id not in by_id or role not in ROLES:
            raise ValueError(f"Unknown sample or role in agent sheet: {key}")
        if key in agents:
            raise ValueError(f"Duplicate agent review for {sample_id}/{role}")
        agents[key] = record
    expected_agents = {(sample_id, role) for sample_id in by_id for role in ROLES}
    if set(agents) != expected_agents:
        raise ValueError("Agent sheet must contain exactly one row for each case and core role")

    result = [json.loads(json.dumps(row)) for row in source_rows]
    for row in result:
        sample_id = row["sample"]["sample_id"]
        ruling_record = rulings[sample_id]
        decision, meta = _decision_fields(ruling_record, f"{sample_id} ruling")
        row["ruling_review"] = meta
        row["approved_ruling"] = decision in {"approve", "correct"}
        row["review_status"] = ("resolved" if decision == "approve" else
                                "overridden" if decision == "correct" else "awaiting_review")
        row["reviewer"] = meta["reviewer"]
        row["approved_at"] = meta["reviewed_at"]
        row["rationale"] = meta["rationale"]
        if decision == "correct":
            raw = _text(ruling_record.get("Corrected ruling JSON"))
            try:
                corrected = json.loads(raw)
            except json.JSONDecodeError as exc:
                raise ValueError(f"{sample_id} ruling: corrected ruling must be valid JSON") from exc
            prompt = json.loads(row["role_targets"]["judge"]["messages"][0]["content"])
            _validate_standalone_judge(corrected, prompt, f"{sample_id} ruling")
            row["ruling_target"] = corrected

        for role in ROLES:
            target = row["role_targets"][role]
            record = agents[(sample_id, role)]
            target_decision, target_meta = _decision_fields(record, f"{sample_id} {role}")
            target["review"] = target_meta
            target["approved"] = target_decision in {"approve", "correct"}
            target["origin"] = "human_reviewed" if target["approved"] else "unreviewed_draft"
            if target_decision == "correct":
                raw = _text(record.get("Corrected target JSON"))
                try:
                    corrected = json.loads(raw)
                except json.JSONDecodeError as exc:
                    raise ValueError(f"{sample_id} {role}: corrected target must be valid JSON") from exc
                prompt = json.loads(target["messages"][0]["content"])
                _target_json({"messages": target["messages"][:-1] + [
                    {"role": "assistant", "content": json.dumps(corrected)}]},
                    role, prompt, f"{sample_id} {role}")
                target["messages"][-1]["content"] = json.dumps(corrected, sort_keys=True)
    return result


def import_workbook(workbook_path, source, output):
    from openpyxl import load_workbook

    destination = outside_repository(output)
    if destination.exists():
        raise FileExistsError(f"Output already exists: {destination}")
    source_rows = [json.loads(line) for line in Path(source).read_text().splitlines() if line.strip()]
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if not {"Rulings", "Agent Targets"}.issubset(workbook.sheetnames):
            raise ValueError("Workbook must include Rulings and Agent Targets worksheets")
        ruling_rows = _records(workbook["Rulings"], {"Sample ID", "Decision"}, "Rulings")
        agent_rows = _records(workbook["Agent Targets"], {"Sample ID", "Agent role", "Decision"}, "Agent Targets")
        result = apply_review_data(source_rows, ruling_rows, agent_rows)
    finally:
        workbook.close()
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text("".join(json.dumps(row, sort_keys=True) + "\n" for row in result))
    approved_rulings = sum(row["approved_ruling"] for row in result)
    approved_targets = sum(target["approved"] for row in result for target in row["role_targets"].values())
    return {"cases": len(result), "approved_rulings": approved_rulings,
            "approved_role_targets": approved_targets, "output": str(destination)}


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True)
    parser.add_argument("--source", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    print(json.dumps(import_workbook(args.workbook, args.source, args.output), indent=2))
